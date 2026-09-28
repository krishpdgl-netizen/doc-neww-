import io, os, re, logging
from difflib import SequenceMatcher
from typing import Any

from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
import pdfplumber
import pytesseract
from pdf2image import convert_from_bytes
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("diffiq")

MAX_FILE = 50 * 1024 * 1024
SUPPORTED = {"pdf", "docx", "xlsx", "xlsm", "pptx", "txt", "md", "csv"}
OCR_DPI = 240
OCR_MIN_CONF = 30

app = FastAPI(title="DiffIQ Document Comparison API", version="5.0.0")
origins = [x.strip() for x in os.getenv("CORS_ORIGINS", "*").split(",") if x.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=origins if origins != ["*"] else ["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


def ext(name: str) -> str:
    return name.lower().rsplit(".", 1)[-1] if "." in name else ""


def norm(s: Any) -> str:
    s = "" if s is None else str(s)
    s = s.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    s = re.sub(r"[ \t]+", " ", s)
    return s.strip()


def canonical(s: str) -> str:
    """Stable comparison key; deliberately conservative so real text changes remain visible."""
    s = norm(s).casefold()
    s = s.replace("“", '"').replace("”", '"').replace("‘", "'").replace("’", "'")
    # PDF extraction sometimes inserts soft hyphens or zero-width characters.
    s = re.sub(r"[\u00ad\u200b\u200c\u200d]", "", s)
    return s


def word_key(s: str) -> str:
    # Keep numbers, dates, punctuation and symbols meaningful.
    return re.sub(r"\s+", " ", canonical(s))


def pdf_page_key(page: dict) -> str:
    # A compact page fingerprint used only to align pages; no highlighting is based on it.
    words = page.get("words", [])
    vals = [word_key(w.get("text", "")) for w in words if word_key(w.get("text", ""))]
    if not vals:
        return ""
    # Include enough content to distinguish neighboring pages while keeping alignment cheap.
    if len(vals) > 180:
        vals = vals[:90] + vals[-90:]
    return " ".join(vals)


def union_bbox(items):
    boxes = [x.get("bbox") for x in items if isinstance(x, dict) and isinstance(x.get("bbox"), list) and len(x["bbox"]) >= 4]
    if not boxes:
        return None
    return [
        round(min(b[0] for b in boxes), 2),
        round(min(b[1] for b in boxes), 2),
        round(max(b[2] for b in boxes), 2),
        round(max(b[3] for b in boxes), 2),
    ]


def make_word(text, x0, y0, x1, y1, page_no, source):
    return {
        "text": norm(text),
        "bbox": [round(float(x0), 2), round(float(y0), 2), round(float(x1), 2), round(float(y1), 2)],
        "page": page_no,
        "source": source,
    }


def group_words_into_lines(words, page_no, source):
    if not words:
        return []
    rows = []
    for w in words:
        top = float(w["y0"])
        best = None
        best_delta = 999
        for row in rows:
            row_y = sum(float(x["y0"]) for x in row) / len(row)
            delta = abs(top - row_y)
            # A little tolerance handles font/subpixel differences without merging separate paragraphs.
            if delta <= 3.5 and delta < best_delta:
                best = row
                best_delta = delta
        if best is None:
            rows.append([w])
        else:
            best.append(w)

    rows.sort(key=lambda r: (min(float(x["y0"]) for x in r), min(float(x["x0"]) for x in r)))
    lines = []
    for row in rows:
        row.sort(key=lambda x: float(x["x0"]))
        line_words = [make_word(w["text"], w["x0"], w["y0"], w["x1"], w["y1"], page_no, source) for w in row if norm(w["text"])]
        if not line_words:
            continue
        lines.append({
            "text": " ".join(w["text"] for w in line_words),
            "page": page_no,
            "bbox": union_bbox(line_words),
            "items": line_words,
            "kind": "line",
            "source": source,
        })
    return lines


def extract_native_page(page, page_no):
    raw = page.extract_words(
        x_tolerance=2.5,
        y_tolerance=3,
        keep_blank_chars=False,
        use_text_flow=False,
        extra_attrs=["size"],
    )
    words = []
    for w in raw or []:
        text = norm(w.get("text", ""))
        if not text:
            continue
        words.append({
            "text": text,
            "x0": float(w["x0"]),
            "y0": float(w["top"]),
            "x1": float(w["x1"]),
            "y1": float(w["bottom"]),
        })
    return group_words_into_lines(words, page_no, "text")


def extract_ocr_page(data, page_no, page_width, page_height):
    img = convert_from_bytes(data, dpi=OCR_DPI, first_page=page_no, last_page=page_no, fmt="PNG")[0]
    sx, sy = page_width / img.width, page_height / img.height
    d = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT, config="--psm 6")
    rows = {}
    for i, raw in enumerate(d["text"]):
        text = norm(raw)
        if not text:
            continue
        try:
            conf = float(d["conf"][i])
        except Exception:
            conf = 0
        if conf < OCR_MIN_CONF:
            continue
        x, y, w, h = d["left"][i], d["top"][i], d["width"][i], d["height"][i]
        key = (int(d["block_num"][i]), int(d["par_num"][i]), int(d["line_num"][i]))
        rows.setdefault(key, []).append({
            "text": text,
            "x0": x * sx,
            "y0": y * sy,
            "x1": (x + w) * sx,
            "y1": (y + h) * sy,
        })
    lines = []
    for words in rows.values():
        lines.extend(group_words_into_lines(words, page_no, "ocr"))
    lines.sort(key=lambda b: (b["page"], b["bbox"][1] if b.get("bbox") else 0, b["bbox"][0] if b.get("bbox") else 0))
    return lines


def extract_pdf(data: bytes):
    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            native_lines = extract_native_page(page, pno)
            native_words = [w for line in native_lines for w in line["items"]]
            native_chars = sum(len(w["text"]) for w in native_words)

            # Native text is preferred. OCR is used only when the page is genuinely image/scanned-like.
            if native_chars >= 20 and len(native_words) >= 3:
                lines = native_lines
                source = "text"
            else:
                try:
                    lines = extract_ocr_page(data, pno, page.width, page.height)
                    source = "ocr" if lines else "empty"
                except Exception:
                    log.exception("OCR failed on page %s", pno)
                    lines = []
                    source = "empty"

            words = [w for line in lines for w in line["items"]]
            pages.append({
                "page": pno,
                "width": float(page.width),
                "height": float(page.height),
                "source": source,
                "blocks": lines,
                "words": words,
            })

    return {
        "format": "pdf",
        "pageCount": len(pages),
        "ocrPages": sum(1 for p in pages if p["source"] == "ocr"),
        "pages": pages,
        "blocks": [b for p in pages for b in p["blocks"]],
        "words": [w for p in pages for w in p["words"]],
    }


def extract_docx(data: bytes):
    doc = Document(io.BytesIO(data))
    blocks = []
    for i, p in enumerate(doc.paragraphs):
        t = norm(p.text)
        if t:
            blocks.append({"text": t, "index": i, "kind": "heading" if p.style and p.style.name.lower().startswith("heading") else "paragraph"})
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            vals = [norm(c.text) for c in row.cells]
            if any(vals):
                blocks.append({"text": " | ".join(vals), "table": ti, "row": ri, "kind": "table-row"})
    return {"format": "docx", "pageCount": None, "pages": [], "blocks": blocks}


def extract_xlsx(data: bytes):
    wb = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    blocks, sheets = [], []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows():
            vals = ["" if c.value is None else str(c.value) for c in row]
            if any(v.strip() for v in vals):
                rows.append({"text": "\t".join(vals).rstrip("\t"), "sheet": ws.title, "row": row[0].row, "kind": "sheet-row"})
        sheets.append({"name": ws.title, "rows": len(rows)})
        blocks.extend(rows)
    return {"format": "xlsx", "pageCount": None, "sheets": sheets, "pages": [], "blocks": blocks}


def extract_pptx(data: bytes):
    prs = Presentation(io.BytesIO(data))
    slides, blocks = [], []
    for si, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and norm(shape.text):
                texts.append(norm(shape.text))
        text = "\n".join(texts)
        slides.append({"slide": si, "text": text})
        if text:
            blocks.append({"text": text, "slide": si, "kind": "slide"})
    return {"format": "pptx", "pageCount": len(prs.slides), "slides": slides, "pages": [], "blocks": blocks}


def extract_plain(data: bytes, extension: str):
    text = data.decode("utf-8-sig", errors="replace")
    lines = [norm(x) for x in text.splitlines() if norm(x)]
    return {"format": extension, "pageCount": None, "pages": [], "blocks": [{"text": x, "line": i + 1, "kind": "line"} for i, x in enumerate(lines)]}


def extract(data, filename):
    e = ext(filename)
    if e not in SUPPORTED:
        raise HTTPException(400, f"Unsupported file type: .{e}. Supported: {', '.join(sorted(SUPPORTED))}")
    if e == "pdf":
        return extract_pdf(data)
    if e == "docx":
        return extract_docx(data)
    if e in {"xlsx", "xlsm"}:
        return extract_xlsx(data)
    if e == "pptx":
        return extract_pptx(data)
    return extract_plain(data, e)


def token_list(s: str):
    return re.findall(r"\w+(?:[.%$₹€£-]\w+)*|[^\w\s]", norm(s), flags=re.UNICODE)


def token_key(s: str):
    return re.sub(r"\s+", " ", canonical(s))


def make_word_block(words):
    items = [{"str": w["text"], "bbox": w["bbox"], "page": w["page"]} for w in words]
    return {
        "text": " ".join(w["text"] for w in words),
        "page": words[0]["page"] if words else None,
        "bbox": union_bbox(items),
        "items": items,
        "diffItems": items,
        "kind": "pdf-word-diff",
        "source": words[0].get("source", "pdf") if words else "pdf",
    }


def line_key(line):
    return [word_key(w["text"]) for w in line.get("items", []) if word_key(w["text"])]


def line_similarity(a, b):
    ka, kb = line_key(a), line_key(b)
    if not ka and not kb:
        return 1.0
    return SequenceMatcher(None, ka, kb, autojunk=False).ratio()


def word_level_change(left_line, right_line):
    lw = left_line.get("items", []) if left_line else []
    rw = right_line.get("items", []) if right_line else []
    ka = [word_key(w["text"]) for w in lw]
    kb = [word_key(w["text"]) for w in rw]
    sm = SequenceMatcher(None, ka, kb, autojunk=False)
    changed_left, changed_right, tokens = [], [], []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        lseg, rseg = lw[i1:i2], rw[j1:j2]
        if lseg:
            changed_left.extend(lseg)
        if rseg:
            changed_right.extend(rseg)
        if tag == "delete":
            tokens.append({"type": "removed", "left": " ".join(w["text"] for w in lseg), "right": ""})
        elif tag == "insert":
            tokens.append({"type": "added", "left": "", "right": " ".join(w["text"] for w in rseg)})
        else:
            tokens.append({"type": "modified", "left": " ".join(w["text"] for w in lseg), "right": " ".join(w["text"] for w in rseg)})

    # If the line-level match is so poor that every word differs, this is still a valid line change;
    # highlight only the actual words in the changed line, never an entire page.
    if not changed_left and not changed_right:
        return None
    return {
        "type": "modified" if changed_left and changed_right else ("removed" if changed_left else "added"),
        "left": make_word_block(changed_left) if changed_left else None,
        "right": make_word_block(changed_right) if changed_right else None,
        "tokens": tokens,
    }


def page_alignment(left_pages, right_pages):
    # Align pages by content fingerprint, allowing inserted/deleted pages without shifting all later pages.
    A = [pdf_page_key(p) for p in left_pages]
    B = [pdf_page_key(p) for p in right_pages]
    sm = SequenceMatcher(None, A, B, autojunk=False)
    pairs = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            pairs.extend((i, j, "equal") for i, j in zip(range(i1, i2), range(j1, j2)))
        elif tag == "replace":
            n = min(i2 - i1, j2 - j1)
            # Pair likely corresponding pages; don't create cross-page giant diffs.
            for k in range(n):
                pairs.append((i1 + k, j1 + k, "replace"))
            for i in range(i1 + n, i2):
                pairs.append((i, None, "delete"))
            for j in range(j1 + n, j2):
                pairs.append((None, j, "insert"))
        elif tag == "delete":
            pairs.extend((i, None, "delete") for i in range(i1, i2))
        elif tag == "insert":
            pairs.extend((None, j, "insert") for j in range(j1, j2))
    return pairs


def pdf_diff(left, right):
    lp, rp = left["pages"], right["pages"]
    changes = []
    same_word_count = 0
    total_left = total_right = 0

    for li, ri, mode in page_alignment(lp, rp):
        lpage = lp[li] if li is not None else None
        rpage = rp[ri] if ri is not None else None
        if lpage:
            total_left += len(lpage.get("words", []))
        if rpage:
            total_right += len(rpage.get("words", []))

        if mode == "delete":
            words = lpage.get("words", []) if lpage else []
            if words:
                changes.append({"type": "removed", "left": make_word_block(words), "right": None,
                                "tokens": [{"type": "removed", "left": " ".join(w["text"] for w in words), "right": ""}]})
            continue
        if mode == "insert":
            words = rpage.get("words", []) if rpage else []
            if words:
                changes.append({"type": "added", "left": None, "right": make_word_block(words),
                                "tokens": [{"type": "added", "left": "", "right": " ".join(w["text"] for w in words)}]})
            continue
        if not lpage or not rpage:
            continue

        ll = lpage.get("blocks", [])
        rr = rpage.get("blocks", [])
        A = [tuple(line_key(x)) for x in ll]
        B = [tuple(line_key(x)) for x in rr]
        # Compare line signatures. A line signature is a list of words, so formatting/spacing changes do not explode the diff.
        sm = SequenceMatcher(None, A, B, autojunk=False)
        for tag, i1, i2, j1, j2 in sm.get_opcodes():
            if tag == "equal":
                same_word_count += sum(len(ll[k].get("items", [])) for k in range(i1, i2))
                continue
            if tag == "delete":
                for k in range(i1, i2):
                    words = ll[k].get("items", [])
                    if words:
                        changes.append({"type": "removed", "left": make_word_block(words), "right": None,
                                        "tokens": [{"type": "removed", "left": ll[k]["text"], "right": ""}]})
                continue
            if tag == "insert":
                for k in range(j1, j2):
                    words = rr[k].get("items", [])
                    if words:
                        changes.append({"type": "added", "left": None, "right": make_word_block(words),
                                        "tokens": [{"type": "added", "left": "", "right": rr[k]["text"]}]})
                continue

            # Replacement: align lines by content, not by position. This is important when
            # Word->PDF conversion re-wraps a paragraph after one word changes.
            lblock = ll[i1:i2]
            rblock = rr[j1:j2]
            LA = [tuple(line_key(x)) for x in lblock]
            RB = [tuple(line_key(x)) for x in rblock]
            lsm = SequenceMatcher(None, LA, RB, autojunk=False)
            for ltag, a1, a2, b1, b2 in lsm.get_opcodes():
                if ltag == "equal":
                    same_word_count += sum(len(lblock[k].get("items", [])) for k in range(a1, a2))
                elif ltag == "delete":
                    for k in range(a1, a2):
                        words = lblock[k].get("items", [])
                        if words:
                            changes.append({"type": "removed", "left": make_word_block(words), "right": None,
                                            "tokens": [{"type": "removed", "left": lblock[k]["text"], "right": ""}]})
                elif ltag == "insert":
                    for k in range(b1, b2):
                        words = rblock[k].get("items", [])
                        if words:
                            changes.append({"type": "added", "left": None, "right": make_word_block(words),
                                            "tokens": [{"type": "added", "left": "", "right": rblock[k]["text"]}]})
                else:
                    ln2, rn2 = a2 - a1, b2 - b1
                    # If wrapping differs, flatten this small replacement region so common
                    # words can anchor the diff across line boundaries.
                    if ln2 != rn2 and (ln2 <= 3 and rn2 <= 3):
                        lw = [w for k in range(a1, a2) for w in lblock[k].get("items", [])]
                        rw = [w for k in range(b1, b2) for w in rblock[k].get("items", [])]
                        ka, kb = [word_key(w["text"]) for w in lw], [word_key(w["text"]) for w in rw]
                        wsm = SequenceMatcher(None, ka, kb, autojunk=False)
                        for wtag, wi1, wi2, wj1, wj2 in wsm.get_opcodes():
                            if wtag == "equal":
                                same_word_count += wi2 - wi1
                                continue
                            lwords, rwords = lw[wi1:wi2], rw[wj1:wj2]
                            if wtag == "delete":
                                changes.append({"type":"removed","left":make_word_block(lwords),"right":None,"tokens":[{"type":"removed","left":" ".join(w["text"] for w in lwords),"right":""}]})
                            elif wtag == "insert":
                                changes.append({"type":"added","left":None,"right":make_word_block(rwords),"tokens":[{"type":"added","left":"","right":" ".join(w["text"] for w in rwords)}]})
                            else:
                                changes.append({"type":"modified","left":make_word_block(lwords),"right":make_word_block(rwords),"tokens":[{"type":"modified","left":" ".join(w["text"] for w in lwords),"right":" ".join(w["text"] for w in rwords)}]})
                    else:
                        n2 = min(ln2, rn2)
                        for k in range(n2):
                            c = word_level_change(lblock[a1+k], rblock[b1+k])
                            if c:
                                changes.append(c)
                            else:
                                same_word_count += min(len(lblock[a1+k].get("items", [])), len(rblock[b1+k].get("items", [])))
                        for k in range(a1+n2, a2):
                            words = lblock[k].get("items", [])
                            if words:
                                changes.append({"type":"removed","left":make_word_block(words),"right":None,"tokens":[{"type":"removed","left":lblock[k]["text"],"right":""}]})
                        for k in range(b1+n2, b2):
                            words = rblock[k].get("items", [])
                            if words:
                                changes.append({"type":"added","left":None,"right":make_word_block(words),"tokens":[{"type":"added","left":"","right":rblock[k]["text"]}]})

    summary = {"same": same_word_count, "added": 0, "removed": 0, "modified": 0}
    for c in changes:
        summary[c["type"]] += 1
    denominator = max(total_left, total_right, 1)
    changed_word_estimate = 0
    for c in changes:
        for side in ("left", "right"):
            b = c.get(side)
            if b and isinstance(b.get("items"), list):
                changed_word_estimate += len(b["items"])
    similarity = max(0.0, min(100.0, (1 - changed_word_estimate / denominator) * 100))
    return {
        "summary": summary,
        "similarity": round(similarity, 2),
        "changes": changes,
        "left": {k: v for k, v in left.items() if k not in {"blocks", "words"}},
        "right": {k: v for k, v in right.items() if k not in {"blocks", "words"}},
        "engine": "DiffIQ PDF page/line/word coordinate comparison v5",
    }


def block_diff(a, b):
    A = [x["text"] for x in a["blocks"]]
    B = [x["text"] for x in b["blocks"]]
    sm = SequenceMatcher(None, A, B, autojunk=False)
    changes = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            for k in range(i2 - i1):
                changes.append({"type": "same", "left": a["blocks"][i1 + k], "right": b["blocks"][j1 + k]})
        elif tag == "delete":
            for k in range(i1, i2):
                changes.append({"type": "removed", "left": a["blocks"][k], "right": None})
        elif tag == "insert":
            for k in range(j1, j2):
                changes.append({"type": "added", "left": None, "right": b["blocks"][k]})
        else:
            n = max(i2 - i1, j2 - j1)
            for k in range(n):
                l = a["blocks"][i1 + k] if i1 + k < i2 else None
                r = b["blocks"][j1 + k] if j1 + k < j2 else None
                if l and r:
                    changes.append({"type": "modified", "left": l, "right": r, "tokens": token_diff(l["text"], r["text"])})
                elif l:
                    changes.append({"type": "removed", "left": l, "right": None})
                else:
                    changes.append({"type": "added", "left": None, "right": r})
    return changes


def token_diff(a, b):
    A, B = token_list(a), token_list(b)
    sm = SequenceMatcher(None, [token_key(x) for x in A], [token_key(x) for x in B], autojunk=False)
    out = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            out.append({"type": "same", "left": " ".join(A[i1:i2]), "right": " ".join(B[j1:j2])})
        elif tag == "delete":
            out.append({"type": "removed", "left": " ".join(A[i1:i2]), "right": ""})
        elif tag == "insert":
            out.append({"type": "added", "left": "", "right": " ".join(B[j1:j2])})
        else:
            out.append({"type": "modified", "left": " ".join(A[i1:i2]), "right": " ".join(B[j1:j2])})
    return out


def compare(left, right):
    if left.get("format") == "pdf" and right.get("format") == "pdf":
        return pdf_diff(left, right)
    changes = block_diff(left, right)
    summary = {"same": 0, "added": 0, "removed": 0, "modified": 0}
    for c in changes:
        summary[c["type"]] += 1
    A = [x["text"] for x in left["blocks"]]
    B = [x["text"] for x in right["blocks"]]
    ratio = SequenceMatcher(None, A, B, autojunk=False).ratio() if A or B else 1
    return {
        "summary": summary,
        "similarity": round(ratio * 100, 2),
        "changes": changes,
        "left": {k: v for k, v in left.items() if k != "blocks"},
        "right": {k: v for k, v in right.items() if k != "blocks"},
        "engine": "DiffIQ structural + token comparison v5",
    }


@app.get("/")
def root():
    return {"service": "DiffIQ API", "version": "5.0.0", "status": "ok", "formats": sorted(SUPPORTED)}


@app.get("/health")
def health():
    return {"status": "ok", "version": "5.0.0"}


@app.post("/compare")
async def compare_endpoint(left: UploadFile = File(...), right: UploadFile = File(...)):
    if ext(left.filename) != ext(right.filename):
        raise HTTPException(400, "For reliable structural comparison, both documents must have the same file type.")
    lb = await left.read()
    rb = await right.read()
    if len(lb) > MAX_FILE or len(rb) > MAX_FILE:
        raise HTTPException(413, "Each file must be 50 MB or smaller.")
    if not lb or not rb:
        raise HTTPException(400, "Both files are required.")
    try:
        le = extract(lb, left.filename)
        re_ = extract(rb, right.filename)
        result = compare(le, re_)
        result["files"] = {"left": left.filename, "right": right.filename}
        return JSONResponse(content=result)
    except HTTPException:
        raise
    except Exception as e:
        log.exception("comparison failed")
        raise HTTPException(500, f"Comparison failed: {type(e).__name__}: {e}")


@app.post("/extract")
async def extract_endpoint(file: UploadFile = File(...)):
    data = await file.read()
    if len(data) > MAX_FILE:
        raise HTTPException(413, "File too large (max 50 MB).")
    try:
        return JSONResponse(content=extract(data, file.filename))
    except HTTPException:
        raise
    except Exception as e:
        log.exception("extraction failed")
        raise HTTPException(500, f"Extraction failed: {type(e).__name__}: {e}")
