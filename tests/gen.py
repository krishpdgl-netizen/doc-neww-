import os, io, random, subprocess, glob
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
import docx, openpyxl, img2pdf
from PIL import Image, ImageFilter
import numpy as np
random.seed(7); np.random.seed(7)
C='corpus/'
heads=["SERVICE AGREEMENT"]
def base(v):
    days = "30" if v==1 else "45"
    date = "01/04/2026" if v==1 else "01/07/2026"
    amt  = "Rs. 5,00,000" if v==1 else "Rs. 6,00,000"
    verb = "shall" if v==1 else "may"
    P=[
    "This Service Agreement is entered into on %s between Acme Industries Private Limited (the Client) and Zenith Technologies LLP (the Vendor)."%date,
    "1. Scope of Work. The Vendor will design, develop and deliver a document management platform in accordance with the specifications set out in Schedule A, including integration with the Client's existing accounting systems and quarterly training for staff.",
    "2. Payment. The Client will pay the Vendor a fixed fee of %s, invoiced in three equal instalments. Each invoice must be settled within %s days of receipt, failing which interest accrues at one percent per month."%(amt,days),
    "3. Confidentiality. Each party agrees to keep the other party's proprietary information strictly confidential and not to disclose it to any third party without prior written consent, during the term and for five years thereafter.",
    "4. Termination. Either party %s terminate this Agreement by giving ninety days written notice to the other party, and the Vendor must hand over all work product completed up to the termination date."%verb,
    "5. Intellectual Property. All deliverables created under this Agreement will be owned by the Client upon receipt of full payment, while the Vendor retains ownership of its pre-existing tools and libraries.",
    "6. Governing Law. This Agreement is governed by the laws of India and the courts at Mumbai shall have exclusive jurisdiction over any dispute arising out of it.",
    ]
    if v==1: P.insert(6,"7. Non-Compete. During the term and for twelve months after, the Vendor will not build a competing product for any direct competitor of the Client within India.")
    if v==2:
        P.insert(6,"7. Force Majeure. Neither party is liable for delay caused by events beyond its reasonable control, including natural disasters, war, epidemics or government action, provided prompt notice is given.")
        c=P.pop(3); P.append(c)   # move Confidentiality to the end
    return P
def pdf(path,P):
    ss=getSampleStyleSheet(); st=ss['BodyText']; st.fontSize=11; st.leading=15; st.spaceAfter=10
    doc=SimpleDocTemplate(path,pagesize=A4,leftMargin=60,rightMargin=60,topMargin=70,bottomMargin=70)
    doc.build([Paragraph("<b>SERVICE AGREEMENT</b>",ss['Title'])]+[Paragraph(p.replace('&','&amp;'),st) for p in P])
for v in (1,2):
    P=base(v)
    open(C+f'contract_v{v}.txt','w').write("SERVICE AGREEMENT\n\n"+"\n\n".join(P))
    d=docx.Document(); d.add_heading("SERVICE AGREEMENT",1)
    for p in P: d.add_paragraph(p)
    d.save(C+f'contract_v{v}.docx')
    open(C+f'contract_v{v}.html','w').write("<html><head><style>p{color:red}</style></head><body><h1>SERVICE AGREEMENT</h1>"+"".join(f"<p>{p}</p>" for p in P)+"<script>alert(1)</script></body></html>")
    pdf(C+f'contract_v{v}_native.pdf',P)
    # csv/xlsx price list
    rows=[["Item","Qty","Unit price","Total"],["Server","2","1,20,000" if v==1 else "1,25,000","2,40,000" if v==1 else "2,50,000"],["Licence, annual","1","50,000","50,000"],["Support, \"premium\"","12","8,000","96,000"]]
    if v==2: rows.append(["Backup storage","1","15,000","15,000"])
    import csv; csv.writer(open(C+f'prices_v{v}.csv','w',newline='')).writerows(rows)
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Prices"
    for r in rows: ws.append(r)
    wb.save(C+f'prices_v{v}.xlsx')
# scanned versions: rasterise natives, degrade differently, rebuild as image-only PDFs
def scan(native,out,angle,shift,noise,blur,q):
    tmp=out+'_pg'; subprocess.run(['pdftoppm','-r','200','-png',native,tmp],check=True)
    jpgs=[]
    for f in sorted(glob.glob(tmp+'*.png')):
        im=Image.open(f).convert('L')
        bg=Image.new('L',im.size,235); bg.paste(im,(shift,shift//2)); im=bg.rotate(angle,resample=Image.BICUBIC,fillcolor=228)
        a=np.asarray(im,dtype=np.float32)+np.random.normal(0,noise,im.size[::-1])
        # uneven lighting gradient like a phone/flatbed scan
        a=a*(0.93+0.07*np.linspace(0,1,a.shape[1])[None,:])
        im=Image.fromarray(np.clip(a,0,255).astype(np.uint8)).filter(ImageFilter.GaussianBlur(blur))
        jf=f.replace('.png','.jpg'); im.save(jf,quality=q); jpgs.append(jf); os.remove(f)
    open(out,'wb').write(img2pdf.convert(jpgs))
    for j in jpgs: os.remove(j)
scan(C+'contract_v1_native.pdf',C+'contract_v1_scan.pdf',0.7,10,7,0.6,70)
scan(C+'contract_v2_native.pdf',C+'contract_v2_scan.pdf',-0.5,22,11,0.8,55)
print(sorted(os.listdir(C)))
